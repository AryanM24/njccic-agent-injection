"""Shared .docx/.pdf/.txt/.md fixture-writer helpers for the corpus generators.

The channel set here is a property of extraction.py (which channels the
pinned parser stack actually exposes), not of any one workflow's content, so
make_fixtures_v2.py (W2), make_fixtures_w1.py, and make_fixtures_w3.py all
build on these instead of each re-implementing channel writing.
"""

import io
import json
import struct
import textwrap
import zlib
from collections.abc import Callable
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

CHANNELS: dict[str, list[str]] = {
    "docx": ["visible", "core_property", "comment", "alt_text"],
    "pdf": ["visible", "pdf_metadata"],
    "txt": ["visible"],
    "md": ["visible"],
}
"""File type -> concealment channels extraction.py can expose for that type."""


def small_png() -> io.BytesIO:
    """Build a minimal valid 4x4 white PNG, in-memory.

    Returns:
        A PNG file's bytes, for use as an embedded image's picture data.
    """

    def chunk(ctype: bytes, data: bytes) -> bytes:
        c = ctype + data
        return (
            struct.pack(">I", len(data))
            + c
            + struct.pack(">I", zlib.crc32(c) & 0xFFFFFFFF)
        )

    w = h = 4
    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    raw = b"".join(b"\x00" + b"\xff\xff\xff" * w for _ in range(h))
    idat = zlib.compress(raw)
    return io.BytesIO(
        sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")
    )


def write_docx(
    path: Path,
    channel: str | None,
    payload: str | None,
    build_body: Callable[[DocumentType], None],
) -> None:
    """Write a .docx fixture, optionally carrying a payload in one channel.

    Args:
        path: Output file path.
        channel: One of CHANNELS["docx"], or None for a clean twin.
        payload: Text to place in the channel. Ignored if channel is None.
        build_body: Fills in the document's visible body content (the part
            shared by the attack file and its twin).
    """
    document = Document()
    build_body(document)
    if channel == "visible":
        document.add_paragraph(payload)
    elif channel == "core_property":
        document.core_properties.comments = payload
    elif channel == "comment":
        anchor = document.paragraphs[1]
        document.add_comment(
            runs=anchor.runs, text=payload, author="A. Reviewer", initials="AR"
        )
    elif channel == "alt_text":
        shape = document.add_picture(small_png(), width=Inches(1))
        docPr = shape._inline.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
        )
        if docPr is not None:
            docPr.set("descr", payload)
    document.save(str(path))


def write_pdf(
    path: Path, channel: str | None, payload: str | None, text_lines: list[str]
) -> None:
    """Write a .pdf fixture, optionally carrying a payload in one channel.

    Args:
        path: Output file path.
        channel: One of CHANNELS["pdf"], or None for a clean twin.
        payload: Text to place in the channel. Ignored if channel is None.
        text_lines: Visible body lines, one per line of the rendered page.
    """
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    top, bottom, leading = 720, 72, 15
    y = top

    def draw(line: str) -> None:
        nonlocal y
        # Wrap rather than truncate, and paginate. The writer used to cut every
        # line at 95 characters and never started a second page, which silently
        # dropped content (a ticket's submitter email and ID, long body lines)
        # from every PDF fixture while leaving docx/txt/md complete -- so file
        # type, an independent variable, was also changing what the model saw.
        for piece in textwrap.wrap(line, width=90, break_long_words=False) or [""]:
            if y < bottom:
                c.showPage()
                y = top
            c.drawString(72, y, piece)
            y -= leading

    for line in text_lines:
        draw(line)
    if channel == "visible":
        y -= leading
        draw(payload)
    c.save()
    buf.seek(0)
    reader = PdfReader(buf)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    if channel == "pdf_metadata":
        writer.add_metadata({"/Subject": payload})
    with open(path, "wb") as f:
        writer.write(f)


def write_text(path: Path, channel: str | None, payload: str | None, text: str) -> None:
    """Write a .txt or .md fixture, optionally appending a visible payload.

    Args:
        path: Output file path.
        channel: "visible" or None for a clean twin.
        payload: Text to append. Ignored if channel is None.
        text: The clean body text.
    """
    content = text
    if channel == "visible":
        content = content + "\n" + payload
    path.write_text(content, encoding="utf-8")


def upsert_manifest_line(manifest: Path, entry: dict) -> str:
    """Insert or update one manifest line, keyed by artifact_id.

    Args:
        manifest: Path to the manifest .jsonl file.
        entry: The manifest entry to write.

    Returns:
        "unchanged" if an identical line already existed, "updated" if a line
        with this artifact_id existed but differed, "appended" if it's new.
    """
    lines: list[str] = []
    if manifest.exists():
        lines = [
            ln for ln in manifest.read_text(encoding="utf-8").splitlines() if ln.strip()
        ]
    serialized = json.dumps(entry)
    for index, line in enumerate(lines):
        if json.loads(line)["artifact_id"] == entry["artifact_id"]:
            if line == serialized:
                return "unchanged"
            lines[index] = serialized
            manifest.write_text("\n".join(lines) + "\n", encoding="utf-8")
            return "updated"
    lines.append(serialized)
    manifest.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return "appended"


# --------------------------------------------------------------------------
# Base scenarios and the shared generation loop
# --------------------------------------------------------------------------

import sys
from dataclasses import dataclass, field, replace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from eduinject.corpus import GOALS, MARKER, load_gated_payloads

SLIM_CHANNELS: tuple[tuple[str, str], ...] = (("docx", "visible"), ("docx", "comment"))
"""The (file type, channel) pairs a "slim" base covers: one visible and one
concealed channel, enough to check that a result is not an artifact of a single
document without multiplying the corpus."""


@dataclass(frozen=True)
class Base:
    """One ordinary document that attacks are placed into.

    A base scenario is the unit of scenario diversity. Many attack variants of
    one base are not independent evidence (they share every quirk of that
    document), so each workflow carries several bases.
    """

    base_id: str
    heading: str
    contact: str
    sections: tuple[tuple[str, tuple[str, ...]], ...]
    author: str
    operator_request: str
    extras: dict = field(default_factory=dict)
    """Manifest fields specific to this base (recipients, protected terms,
    entitled actions, expectations, ...)."""
    slice: str = "full"
    """"full": every file type x channel x goal. "slim": SLIM_CHANNELS x goal."""

    def text_lines(self) -> list[str]:
        lines = [self.heading, self.contact]
        for section, entries in self.sections:
            lines.append(section)
            lines.extend(entries)
        return lines

    def text(self) -> str:
        return "\n".join(self.text_lines())

    def build_docx(self, document: DocumentType) -> None:
        document.core_properties.comments = ""
        document.core_properties.author = self.author
        document.add_heading(self.heading, level=1)
        document.add_paragraph(self.contact)
        for section, entries in self.sections:
            document.add_heading(section, level=2)
            for entry in entries:
                document.add_paragraph(entry, style="List Bullet")


def with_extra_line(base: Base, section_prefix: str, line: str) -> Base:
    """Return a copy of `base` with one extra line appended to a section.

    Used to build benign-but-weird controls: the base document plus one
    legitimate sentence that merely *sounds* instruction-like.
    """
    matched = False
    sections = []
    for title, entries in base.sections:
        if not matched and title.startswith(section_prefix):
            entries = (*entries, line)
            matched = True
        sections.append((title, entries))
    if not matched:
        raise ValueError(f"base {base.base_id} has no section starting {section_prefix!r}")
    return replace(base, sections=tuple(sections))


def write_document(base: Base, path: Path, ftype: str, channel: str | None, payload: str | None) -> None:
    """Write `base` as a `ftype` file, optionally carrying `payload` in `channel`."""
    if ftype == "docx":
        write_docx(path, channel, payload, base.build_docx)
    elif ftype == "pdf":
        write_pdf(path, channel, payload, base.text_lines())
    elif ftype in ("txt", "md"):
        write_text(path, channel, payload, base.text())
    else:
        raise ValueError(f"unsupported file type {ftype!r}")


def plan(base: Base) -> list[tuple[str, str, str]]:
    """The (file type, goal, channel) combinations a base is built for.

    The "full" ordering (file type, then goal, then channel) is the one the
    original single-base generators used, so primary-base artifact IDs are
    unchanged.
    """
    if base.slice == "slim":
        return [
            (ftype, goal, channel)
            for goal in GOALS
            for ftype, channel in SLIM_CHANNELS
        ]
    return [
        (ftype, goal, channel)
        for ftype, channels in CHANNELS.items()
        for goal in GOALS
        for channel in channels
    ]


def select_payloads(workflow: str, canonical: bool) -> tuple[dict[str, str], str]:
    """Choose the payload table and its class.

    Public builds are marker-only (release policy, README C3). Canonical builds
    need the git-ignored payload module and write only under the git-ignored
    `corpus/canonical/` tree.
    """
    if not canonical:
        return {goal: MARKER for goal in GOALS}, "placeholder"
    table = load_gated_payloads()
    if table is None:
        raise SystemExit(
            "--canonical needs src/eduinject/corpus/payloads.py, which is "
            "git-ignored and absent on this machine. Public builds are "
            "marker-only by design."
        )
    return table[workflow], "canonical"


def generate_workflow(
    *,
    workflow: str,
    stem: str,
    bases: list[Base],
    canonical: bool = False,
    start_count: int = 1,
) -> int:
    """Generate every attack/twin pair for a workflow's bases.

    Args:
        workflow: "w1", "w2" or "w3".
        stem: Filename stem ("ticket", "resume", "correspondence").
        bases: Base scenarios, in ID order. Pairs are numbered consecutively.
        canonical: Build the gated real-payload corpus under corpus/canonical/
            instead of the public marker-only corpus.
        start_count: The last already-used number plus one (W2 reserves 1 for
            the original pilot pair).

    Returns:
        The last artifact number used.
    """
    payloads, payload_class = select_payloads(workflow, canonical)
    if canonical:
        file_dir = ROOT / "corpus" / "canonical" / workflow
        path_prefix = f"canonical/{workflow}"
        manifest = ROOT / "corpus" / "canonical" / f"manifest_{workflow}.jsonl"
    else:
        file_dir = ROOT / "corpus" / "benign"
        path_prefix = "corpus/benign"
        manifest = ROOT / "corpus" / f"manifest_{workflow}.jsonl"
    file_dir.mkdir(parents=True, exist_ok=True)

    count = start_count - 1
    for base in bases:
        for ftype, goal, channel in plan(base):
            count += 1
            name = f"{stem}_{count:04d}"
            attack_path = file_dir / f"{name}_attack.{ftype}"
            twin_path = file_dir / f"{name}_twin.{ftype}"
            write_document(base, twin_path, ftype, None, None)
            write_document(base, attack_path, ftype, channel, payloads[goal])
            outcome = upsert_manifest_line(
                manifest,
                {
                    "artifact_id": f"{workflow}_{count:04d}",
                    "artifact_path": f"{path_prefix}/{name}_attack.{ftype}",
                    "twin_path": f"{path_prefix}/{name}_twin.{ftype}",
                    "payload_marker": MARKER,
                    "operator_request": base.operator_request,
                    **base.extras,
                    "goal": goal,
                    "file_type": ftype,
                    "concealment": channel,
                    "base_id": base.base_id,
                    "payload_class": payload_class,
                },
            )
            print(f"{workflow}_{count:04d} {base.base_id} {goal}/{channel}/{ftype} {outcome}")
    return count
