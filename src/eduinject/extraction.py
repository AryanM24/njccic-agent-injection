"""Produce "the text the intake agent sees".

This is deliberately *not* a human-faithful rendering of a document. It
includes channels a human reader would never see — PDF metadata, docx core
properties, real tracked comments, image alt text — because those are exactly
the channels an indirect prompt injection can hide in. What this module emits
defines the attack surface of the whole testbed, which is why the parser
versions behind it are pinned (see `eduinject.config.EXTRACTION_PINS`).
"""

import zipfile
from pathlib import Path
from xml.etree import ElementTree

import docx
from pypdf import PdfReader

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP_NS = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
DOCX_COMMENTS_PART = "word/comments.xml"


def extract(path: str | Path) -> str:
    """Extract the full text an intake agent would see from a document.

    Args:
        path: Path to a `.pdf`, `.docx`, `.txt` or `.md` file.

    Returns:
        The extracted text, including non-visible channels.

    Raises:
        ValueError: If the file suffix is not a supported format.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    elif suffix == ".docx":
        return _extract_docx(path)
    elif suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8")
    else:
        raise ValueError(
            f"unsupported document format '{suffix}' for {path} "
            "(expected .pdf, .docx, .txt or .md)"
        )


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))

    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            lines.append(text)

    metadata = reader.metadata
    if metadata:
        for key, value in metadata.items():
            lines.append(f"[metadata {key}] {value}")

    return "\n".join(lines)


def _extract_docx(path: Path) -> str:
    document = docx.Document(str(path))

    lines: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            lines.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    lines.append(cell.text)

    lines.extend(_docx_core_properties(document))
    lines.extend(_docx_comments(path))
    lines.extend(_docx_alt_text(document))

    return "\n".join(lines)


def _docx_core_properties(document: docx.document.Document) -> list[str]:
    """Channel (i): docx core *properties*.

    Note `core_properties.comments` is the OOXML "Comments" metadata field
    (Word's File > Info > Comments box) — it is unrelated to the tracked
    document comments read by `_docx_comments`.
    """
    core = document.core_properties
    lines: list[str] = []
    for name in ("comments", "subject", "keywords", "category"):
        value = getattr(core, name, None)
        if value:
            lines.append(f"[property {name}] {value}")
    return lines


def _docx_comments(path: Path) -> list[str]:
    """Channel (ii): real tracked document comments.

    These live in `word/comments.xml` inside the .docx zip. Reading the part
    directly keeps extraction independent of python-docx's comment object model
    (`Document.comments` also exists as of python-docx 1.2.0) and makes the
    absent-part case — every benign twin — a normal outcome rather than an error.
    """
    with zipfile.ZipFile(path) as archive:
        try:
            part = archive.read(DOCX_COMMENTS_PART)
        except KeyError:
            return []

    root = ElementTree.fromstring(part)
    lines: list[str] = []
    for comment in root.iter(f"{W_NS}comment"):
        text = "".join(node.text or "" for node in comment.iter(f"{W_NS}t"))
        if text:
            lines.append(f"[comment] {text}")
    return lines


def _docx_alt_text(document: docx.document.Document) -> list[str]:
    """Channel (iii): alt text on inline shapes, where cheaply reachable."""
    lines: list[str] = []
    for shape in document.inline_shapes:
        doc_pr = shape._inline.find(f"{WP_NS}docPr")
        if doc_pr is None:
            continue
        for attribute in ("descr", "name"):
            value = doc_pr.get(attribute)
            if value:
                lines.append(f"[alt-text {attribute}] {value}")
    return lines
